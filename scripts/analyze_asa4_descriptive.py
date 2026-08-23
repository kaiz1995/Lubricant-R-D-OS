from __future__ import annotations
"""ASA#4 dose descriptive analysis (Partial Replay, Gate B).

Scope ruling: DESCRIPTIVE evidence strength only.
- No regression / ANOVA / MODEL_BUILT claims (frozen statistics_input requires >=3
  mixture components; frozen validation framework requires doe_result upstream --
  neither contract fits this historical OFAT dose series, so the engine path is
  correctly unavailable).
- No hypothesis tests: descriptive = tables, rank correlation as effect summary.
- Bench noise values measured at different durations conflate dose with test
  stopping time; both dimensions reported, never merged silently.
Deterministic: fixed input file, sorted iteration, math.fsum only.
"""
import json
import math
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATASET = ROOT / "data" / "real_cases" / "WGO_001" / "normalized" / "asa4_dose_bench.json"
SRC_DIR = Path("D:/临时文件-工作/风电-临时文件/风电导电油&抗电蚀剂开发过程记录")
SOLUBILITY_SRC = SRC_DIR / "Lubemater WT 抗电蚀复合剂开发-20260510.xlsx"


def mean(values):
    return math.fsum(values) / len(values)


def bench_stop_status(rec):
    """Classify how the bench run ended. Ambiguity flagged, not guessed."""
    dur, db = rec["bench_duration_h"], rec["bench_noise_db"]
    if dur >= 9.9:
        return "FULL_10H_PASS" if db <= 65 else "FULL_10H_EXCEED"
    # guardrail is "<=65 dB" (≯65): reaching exactly 65 dB means the limit was hit
    if db >= 65:
        return "EARLY_STOP_AT_THRESHOLD"
    return "AMBIGUOUS_STOP_BELOW_LIMIT(E-07/E-08)"


def dose_table(records):
    by_dose = {}
    for rec in sorted(records, key=lambda r: (r["asa4_dose_fraction"], r["run_reference"])):
        by_dose.setdefault(rec["asa4_dose_fraction"], []).append(rec)
    rows = []
    for dose in sorted(by_dose):
        grp = by_dose[dose]
        dbs = [r["bench_noise_db"] for r in grp]
        durs = [r["bench_duration_h"] for r in grp]
        rows.append({
            "dose_fraction": dose,
            "n": len(grp),
            "noise_db": {"min": min(dbs), "max": max(dbs), "mean": round(mean(dbs), 6)},
            "duration_h": {"min": min(durs), "max": max(durs)},
            "stop_status_counts": _count_statuses(grp),
            "run_references": [r["run_reference"] for r in grp],
        })
    return rows


def _count_statuses(grp):
    counts = {}
    for r in grp:
        s = bench_stop_status(r)
        counts[s] = counts.get(s, 0) + 1
    return counts


def spearman_rho(pairs):
    """Rank correlation coefficient. Effect-size summary only; NO p-value."""
    def ranks(vals):
        order = sorted(range(len(vals)), key=lambda i: vals[i])
        rk = [0.0] * len(vals)
        i = 0
        while i < len(order):
            j = i
            while j + 1 < len(order) and vals[order[j + 1]] == vals[order[i]]:
                j += 1
            avg = (i + j) / 2.0 + 1.0
            for k in range(i, j + 1):
                rk[order[k]] = avg
            i = j + 1
        return rk
    xs, ys = zip(*pairs)
    rx, ry = ranks(list(xs)), ranks(list(ys))
    mx, my = mean(rx), mean(ry)
    num = math.fsum((a - mx) * (b - my) for a, b in zip(rx, ry))
    den = math.sqrt(math.fsum((a - mx) ** 2 for a in rx) * math.fsum((b - my) ** 2 for b in ry))
    return num / den if den else 0.0


def solubility_summary():
    """Water content (ppm) vs ASA#4 dose from ASA#4溶解 sheet. Grouped by prep
    condition; conditions differ so cross-group pooling is forbidden."""
    import openpyxl
    wb = openpyxl.load_workbook(SOLUBILITY_SRC, data_only=True)
    ws = wb["ASA#4溶解"]
    groups = {
        "OPEN_DRY_60C_3D": {"cols": ["C", "D", "E", "F"], "doses": [0.005, 0.01, 0.02, 0.03]},
        "CLOSED_THEN_OPEN_50C": {"cols": ["G", "H"], "doses": [0.005, 0.01]},
        "OPEN_50C_STIR_4H": {"cols": ["I", "J", "K"], "doses": [0.05, 0.08, 0.10]},
    }
    out = []
    for gname, g in groups.items():
        pts = []
        for col, dose in zip(g["cols"], g["doses"]):
            v = ws[f"{col}6"].value
            pts.append({"dose_fraction": dose, "water_ppm": v,
                        "cell": f"ASA#4溶解!{col}6", "sample_date_yymmdd": ws[f"{col}5"].value})
        out.append({"prep_condition_group": gname, "points": pts,
                    "note": "within-group prep identical; across-group comparison blocked"})
    neat = {"asa4_neat_water_ppm": ws["B6"].value, "cell": "ASA#4溶解!B6",
            "drying_history_cells": ["ASA#4溶解!B7", "ASA#4溶解!B9"]}
    return {"groups": out, "neat_asa4": neat}


def main() -> int:
    data = json.loads(DATASET.read_text(encoding="utf-8"))
    records = data["records"]

    enriched = []
    for rec in records:
        r = dict(rec)
        r["stop_status"] = bench_stop_status(rec)
        enriched.append(r)

    table = dose_table(enriched)
    pairs_all = [(r["asa4_dose_fraction"], r["bench_noise_db"]) for r in enriched]
    pairs_full_pass = [(r["asa4_dose_fraction"], r["bench_noise_db"])
                       for r in enriched if r["stop_status"] == "FULL_10H_PASS"]
    durations_pairs = [(r["asa4_dose_fraction"], r["bench_duration_h"]) for r in enriched]

    result = {
        "analysis_id": "WGO001_ASA4_DESCRIPTIVE_V1",
        "evidence_strength": "DESCRIPTIVE",
        "model_built_claim": False,
        "input_dataset_sha256_source": records[0]["provenance"]["source_file_sha256"],
        "n_records": len(records),
        "interpretation_boundaries": {
            "gate_ruling": "Partial Replay Gate PASS/FREEZE",
            "dose_vs_noise": "no clear monotonic dose-noise relationship observed in the current historical sample; this does NOT establish that dose has no effect",
            "dose_vs_runtime": "EXPLORATORY_ASSOCIATION ONLY; recorded runtime embeds an unknown stopping mechanism (E-10 CRITICAL) and must not be read as time-to-failure or endurance life",
            "full_run_ranking": "only 2 records completed 10h; ranking over n=2 is not statistical evidence",
            "solubility_minimum": "DESCRIPTIVE_LOCAL_MINIMUM among tested doses; NOT an OPTIMUM",
            "bench_boundary_65db": "TEMPORARY_INTERPRETATION_NOT_METHOD_VERIFIED; boundary_status=UNRESOLVED pending method document review (E-11)",
        },
        "e10_stop_reason_policy": {
            "status": "CRITICAL",
            "resolution": "OWNER_CONFIRMATION_REQUIRED",
            "default_stop_reason": "UNKNOWN",
            "target_classification": ["COMPLETED", "EVENT_OCCURRED", "RIGHT_CENSORED",
                                      "ADMINISTRATIVE_STOP", "INVALID_RUN", "UNKNOWN"],
            "note": "affects 6/10 records stopped below 65 dB before 10h",
        },
        "e11_boundary_status": {
            "boundary_status": "UNRESOLVED",
            "temporary_interpretation": "65.0 dB counted as threshold hit",
            "prohibition": "must not be written into any Test Method Contract until source method document confirms semantics of the guardrail notation",
        },
        "dose_bench_table": table,
        "rank_correlations": {
            "definition": "Spearman rho; effect summary only, no p-value (descriptive scope)",
            "dose_vs_noise_all_n10": round(spearman_rho(pairs_all), 6),
            "dose_vs_noise_full10h_only_n2": (round(spearman_rho(pairs_full_pass), 6)
                if len(pairs_full_pass) >= 2 else None),
            "dose_vs_duration_all_n10": round(spearman_rho(durations_pairs), 6),
            "interpretation_guard": ("negative dose-vs-duration rho is an EXPLORATORY_ASSOCIATION between dose "
                                     "and RECORDED RUNTIME under unknown stopping rules; runtime != time-to-failure "
                                     "!= verified endurance life (E-10)"),
        },
        "solubility_water_content": solubility_summary(),
        "quality_flags": [
            "noise values compared across different durations conflate dose with stopping time",
            "E-10 CRITICAL: 6/10 records stopped below 65dB before 10h with UNKNOWN reason; OWNER_CONFIRMATION_REQUIRED",
            "E-07/E-08: bench operating parameters and retune semantics unresolved; confound interpretation",
            "no replicates (E-05): dispersion within dose is confounded with time/bench drift",
            "bench method version change lacks formal record (E-04); new-standard subset only",
            "65.0 dB counted as threshold hit is TEMPORARY_INTERPRETATION; method document must confirm boundary semantics (E-11)",
        ],
        "descriptive_findings": [],  # filled below deterministically
    }

    findings = result["descriptive_findings"]
    full_pass = [r for r in enriched if r["stop_status"] == "FULL_10H_PASS"]
    findings.append(
        f"{len(full_pass)}/{len(records)} runs completed 10h with noise <=65dB "
        f"({', '.join(r['run_reference'] + ' @' + str(r['asa4_dose_fraction']) for r in full_pass)}).")
    early = [r for r in enriched if r["stop_status"] == "EARLY_STOP_AT_THRESHOLD"]
    if early:
        doses = sorted({r['asa4_dose_fraction'] for r in early})
        findings.append(f"Stops at/above 65dB (boundary reading is TEMPORARY_INTERPRETATION per E-11) occurred at "
                        f"doses {doses}. The dose-vs-runtime negative association is EXPLORATORY_ASSOCIATION only: "
                        f"recorded runtime embeds an unknown stopping mechanism (E-10), so no endurance conclusion is drawn.")
    amb = [r["run_reference"] for r in enriched if r["stop_status"].startswith("AMBIGUOUS")]
    if amb:
        findings.append(f"Ambiguous early stops below limit: {amb} - HOLD from any performance conclusion.")
    sol = result["solubility_water_content"]["groups"]
    g1 = next(g for g in sol if g["prep_condition_group"] == "OPEN_DRY_60C_3D")
    vals = [p["water_ppm"] for p in g1["points"]]
    findings.append(f"Solubility proxy (identical prep OPEN_DRY_60C_3D): water content {vals} ppm at doses "
                    f"{[p['dose_fraction'] for p in g1['points']]}; 2% shows the observed local minimum among tested "
                    f"doses (DESCRIPTIVE_LOCAL_MINIMUM, not OPTIMUM); 2%-vs-3% difference not verified against "
                    f"method repeatability.")

    report_dir = ROOT / "data" / "real_cases" / "WGO_001" / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "ASA4_DESCRIPTIVE_ANALYSIS.json").write_bytes(
        (json.dumps(result, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))

    md = ["# WGO_001 - ASA#4 Dose Descriptive Analysis (PARTIAL REPLAY)",
          "",
          "> Evidence strength: **DESCRIPTIVE**. No model built. No optimization input.",
          "",
          "| dose | n | noise dB(min/max/mean) | duration h(min/max) | stop statuses |",
          "|------|--:|------------------------|---------------------|---------------|"]
    for row in table:
        md.append(f"| {row['dose_fraction']} | {row['n']} | "
                  f"{row['noise_db']['min']}/{row['noise_db']['max']}/{row['noise_db']['mean']} | "
                  f"{row['duration_h']['min']}/{row['duration_h']['max']} | "
                  f"{json.dumps(row['stop_status_counts'], ensure_ascii=False)} |")
    md += ["", "## Findings"]
    md += [f"- {f}" for f in findings]
    md += ["", "## Quality flags"]
    md += [f"- {f}" for f in result["quality_flags"]]
    md += ["", "## Contract boundary",
           "- Frozen statistics engine not invoked: schema requires >=3 mixture components +",
           "  validation requires doe_result upstream; historical OFAT dose series satisfies neither.",
           "- This is the correct system behavior: real imperfect data -> DESCRIPTIVE evidence, not MODEL_BUILT."]
    (report_dir / "ASA4_DESCRIPTIVE_ANALYSIS.md").write_bytes(
        ("\n".join(md) + "\n").encode("utf-8"))
    print(json.dumps({"report": str(report_dir.relative_to(ROOT)), "findings": len(findings)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
