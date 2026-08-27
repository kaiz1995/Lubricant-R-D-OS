"""Local data layer for the lubricant domain pack.

SQLite via stdlib only. Every row carries an evidence_scope so synthetic
demo data can never be mistaken for physical evidence at query level.
"""

from __future__ import annotations

import csv
import json
import sqlite3
from pathlib import Path

TABLES = ("material", "formula", "test_method", "experiment", "benchmark", "artifact")
SCOPES = ("synthetic", "demo", "physical", "unknown")
CLASSIFICATIONS = ("public", "internal", "confidential_formulation", "customer_confidential")
SCHEMA_VERSION = 2
MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
_FORMULA_PREFIXES = ("=", "+", "-", "@")

_SCHEMA = """
CREATE TABLE IF NOT EXISTS {table} (
    id TEXT PRIMARY KEY,
    name TEXT NOT NULL,
    version INTEGER NOT NULL DEFAULT 1,
    evidence_scope TEXT NOT NULL CHECK (evidence_scope IN ('synthetic','demo','physical','unknown')),
    data_classification TEXT NOT NULL DEFAULT 'internal' CHECK (data_classification IN ('public','internal','confidential_formulation','customer_confidential')),
    payload_json TEXT NOT NULL,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
"""

def _migration_backup_path(path: Path, version: int) -> Path:
    candidate = path.with_name(f"{path.stem}.pre-migrate-v{version}.bak")
    suffix = 1
    while candidate.exists():
        candidate = path.with_name(f"{path.stem}.pre-migrate-v{version}-{suffix}.bak")
        suffix += 1
    return candidate


def _backup_to(conn: sqlite3.Connection, target_path: Path) -> None:
    target = sqlite3.connect(str(target_path))
    try:
        conn.backup(target)
    finally:
        target.close()


def connect(path: str | Path) -> sqlite3.Connection:
    db_path = Path(path)
    existing_db = db_path.exists()
    conn = sqlite3.connect(str(db_path))
    try:
        current_version = schema_version(conn)
        if current_version > SCHEMA_VERSION:
            raise ValueError(f"database schema v{current_version} is newer than supported v{SCHEMA_VERSION}")

        has_user_tables = conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' LIMIT 1"
        ).fetchone() is not None
        if existing_db and has_user_tables and current_version < SCHEMA_VERSION:
            _backup_to(conn, _migration_backup_path(db_path, current_version))

        conn.execute("BEGIN IMMEDIATE")
        for table in TABLES:
            conn.execute(_SCHEMA.format(table=table))
            columns = {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}
            if "data_classification" not in columns:
                conn.execute(
                    f"ALTER TABLE {table} ADD COLUMN data_classification TEXT NOT NULL DEFAULT 'internal' "
                    "CHECK (data_classification IN ('public','internal','confidential_formulation','customer_confidential'))"
                )
        conn.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")
        conn.commit()
        return conn
    except Exception:
        conn.rollback()
        conn.close()
        raise


def schema_version(conn: sqlite3.Connection) -> int:
    return int(conn.execute("PRAGMA user_version").fetchone()[0])


def upsert(conn: sqlite3.Connection, table: str, row: dict) -> None:
    row = _validated_row(table, row)
    conn.execute(
        f"INSERT OR REPLACE INTO {table} (id,name,version,evidence_scope,data_classification,payload_json) VALUES (?,?,?,?,?,?)",
        (row["id"], row["name"], row["version"], row["evidence_scope"], row["data_classification"], json.dumps(row["payload"], ensure_ascii=False)),
    )
    conn.commit()


def get(conn: sqlite3.Connection, table: str, id_: str) -> dict | None:
    cur = conn.execute(f"SELECT id,name,version,evidence_scope,data_classification,payload_json FROM {table} WHERE id=?", (id_,))
    row = cur.fetchone()
    if row is None:
        return None
    return {"id": row[0], "name": row[1], "version": row[2], "evidence_scope": row[3], "data_classification": row[4], "payload": json.loads(row[5])}


def list_rows(conn: sqlite3.Connection, table: str, evidence_scope: str | None = None) -> list[dict]:
    sql = f"SELECT id,name,version,evidence_scope,data_classification,payload_json FROM {table}"
    params: tuple = ()
    if evidence_scope is not None:
        sql += " WHERE evidence_scope=?"
        params = (evidence_scope,)
    return [
        {"id": r[0], "name": r[1], "version": r[2], "evidence_scope": r[3], "data_classification": r[4], "payload": json.loads(r[5])}
        for r in conn.execute(sql, params).fetchall()
    ]


def delete(conn: sqlite3.Connection, table: str, id_: str) -> bool:
    cur = conn.execute(f"DELETE FROM {table} WHERE id=?", (id_,))
    conn.commit()
    return cur.rowcount > 0


_CSV_COLS = ("id", "name", "version", "evidence_scope", "data_classification", "payload")
_REQUIRED_IMPORT_COLS = ("id", "name", "version", "evidence_scope", "payload")


def _validated_row(table: str, row: dict) -> dict:
    if table not in TABLES:
        raise ValueError(f"unknown table: {table}")
    try:
        id_ = row["id"]
        name = row["name"]
        payload = row["payload"]
    except KeyError as exc:
        raise ValueError(f"missing row field: {exc.args[0]}") from exc
    if not isinstance(id_, str) or not id_:
        raise ValueError("id must be a non-empty string")
    if not isinstance(name, str) or not name:
        raise ValueError("name must be a non-empty string")
    scope = row.get("evidence_scope")
    if scope not in SCOPES:
        raise ValueError(f"invalid evidence_scope: {scope}")
    classification = row.get("data_classification", "internal")
    if classification not in CLASSIFICATIONS:
        raise ValueError(f"invalid data_classification: {classification}")
    try:
        version = int(row.get("version", 1))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid version: {row.get('version')}") from exc
    return {"id": id_, "name": name, "version": version, "evidence_scope": scope,
            "data_classification": classification, "payload": payload}


def _safe_export_value(value: object) -> object:
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _safe_import_value(value: object, field: str) -> object:
    if isinstance(value, str):
        if value.startswith(_FORMULA_PREFIXES):
            raise ValueError(f"formula-like value rejected in {field}")
        if len(value) > 1 and value[0] == "'" and value[1] in _FORMULA_PREFIXES:
            return value[1:]
    return value


def _ensure_exportable(rows: list[dict]) -> None:
    if any(row["data_classification"] in ("confidential_formulation", "customer_confidential") for row in rows):
        raise ValueError("refusing export of confidential data: this system provides no encryption or ACL")


def _check_import_file(path: str | Path) -> Path:
    file_path = Path(path)
    if file_path.stat().st_size > MAX_IMPORT_BYTES:
        raise ValueError(f"import file exceeds {MAX_IMPORT_BYTES} byte limit")
    return file_path


def _validate_headers(headers: list[object] | tuple[object, ...] | None) -> list[str]:
    if not headers or len(headers) != len(set(headers)):
        raise ValueError("invalid import headers")
    text_headers = [_safe_import_value(value, "header") for value in headers]
    if not all(isinstance(value, str) for value in text_headers) or not set(_REQUIRED_IMPORT_COLS).issubset(text_headers):
        raise ValueError("missing required import columns")
    return text_headers


def _parse_import_row(table: str, raw: dict, row_number: int) -> dict:
    values = {key: _safe_import_value(value, key) for key, value in raw.items()}
    try:
        payload = json.loads(values["payload"])
    except (KeyError, TypeError, json.JSONDecodeError) as exc:
        raise ValueError(f"invalid JSON payload at row {row_number}") from exc
    return _validated_row(table, {"id": values.get("id"), "name": values.get("name"),
                                  "version": values.get("version"), "evidence_scope": values.get("evidence_scope"),
                                  "data_classification": values.get("data_classification") or "internal", "payload": payload})


def _prepare_import(conn: sqlite3.Connection, table: str, raw_rows: list[tuple[int, dict]], collision: str) -> list[dict]:
    if table not in TABLES:
        raise ValueError(f"unknown table: {table}")
    if collision not in ("reject", "replace"):
        raise ValueError(f"invalid collision policy: {collision}")
    rows: list[dict] = []
    ids: set[str] = set()
    for row_number, raw in raw_rows:
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ValueError(f"import exceeds {MAX_IMPORT_ROWS} row limit")
        row = _parse_import_row(table, raw, row_number)
        if row["id"] in ids:
            raise ValueError(f"duplicate import id: {row['id']}")
        ids.add(row["id"])
        rows.append(row)
    if collision == "reject":
        for id_ in ids:
            if conn.execute(f"SELECT 1 FROM {table} WHERE id=?", (id_,)).fetchone() is not None:
                raise ValueError(f"existing id rejected: {id_}")
    return rows


def _commit_import(conn: sqlite3.Connection, table: str, rows: list[dict], collision: str) -> int:
    statement = "INSERT OR REPLACE" if collision == "replace" else "INSERT"
    with conn:
        for row in rows:
            conn.execute(
                f"{statement} INTO {table} (id,name,version,evidence_scope,data_classification,payload_json) VALUES (?,?,?,?,?,?)",
                (row["id"], row["name"], row["version"], row["evidence_scope"], row["data_classification"], json.dumps(row["payload"], ensure_ascii=False)),
            )
    return len(rows)


def export_csv(conn: sqlite3.Connection, table: str, path: str | Path, *, allow_confidential: bool = False) -> None:
    rows = list_rows(conn, table)
    if not allow_confidential:
        _ensure_exportable(rows)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=_CSV_COLS)
        writer.writeheader()
        for row in rows:
            flat = {k: row[k] for k in _CSV_COLS if k != "payload"}
            flat["payload"] = json.dumps(row["payload"], ensure_ascii=False)
            writer.writerow({key: _safe_export_value(value) for key, value in flat.items()})


def import_csv(conn: sqlite3.Connection, table: str, path: str | Path, *, collision: str = "reject") -> int:
    file_path = _check_import_file(path)
    with open(file_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        _validate_headers(reader.fieldnames)
        raw_rows: list[tuple[int, dict]] = []
        for row_number, raw in enumerate(reader, start=2):
            if not any(value is not None and value != "" for value in raw.values()):
                continue
            if len(raw_rows) >= MAX_IMPORT_ROWS:
                raise ValueError(f"import exceeds {MAX_IMPORT_ROWS} row limit")
            raw_rows.append((row_number, raw))
    return _commit_import(conn, table, _prepare_import(conn, table, raw_rows, collision), collision)


def export_snapshot(conn: sqlite3.Connection, path: str | Path, *, allow_confidential: bool = False) -> None:
    snap = {t: list_rows(conn, t) for t in TABLES}
    if not allow_confidential:
        _ensure_exportable([row for rows in snap.values() for row in rows])
    Path(path).write_text(json.dumps(snap, ensure_ascii=False, indent=2), encoding="utf-8")


def backup(conn: sqlite3.Connection, target_path: str | Path) -> None:
    target = sqlite3.connect(str(target_path))
    try:
        conn.backup(target)
    finally:
        target.close()


def restore(source_path: str | Path, target_path: str | Path) -> None:
    target_path = Path(target_path)
    if target_path.exists():
        raise ValueError(f"target already exists: {target_path}")
    source = sqlite3.connect(str(source_path))
    target = sqlite3.connect(str(target_path))
    try:
        source.backup(target)
    finally:
        source.close()
        target.close()


def _openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX requires openpyxl; install it explicitly to enable XLSX I/O") from exc
    return Workbook, load_workbook


def export_xlsx(conn: sqlite3.Connection, table: str, path: str | Path, *, allow_confidential: bool = False) -> None:
    Workbook, _ = _openpyxl()
    rows = list_rows(conn, table)
    if not allow_confidential:
        _ensure_exportable(rows)
    book = Workbook()
    sheet = book.active
    sheet.title = table
    sheet.append(_CSV_COLS)
    for row in rows:
        values = [json.dumps(row["payload"], ensure_ascii=False) if col == "payload" else row[col] for col in _CSV_COLS]
        sheet.append([_safe_export_value(value) for value in values])
    book.save(path)


def import_xlsx(conn: sqlite3.Connection, table: str, path: str | Path, *, collision: str = "reject") -> int:
    _, load_workbook = _openpyxl()
    book = load_workbook(_check_import_file(path), read_only=True, data_only=False)
    try:
        sheet = book.active
        header_cells = next(sheet.iter_rows(), None)
        if header_cells is None or any(cell.data_type == "f" for cell in header_cells):
            raise ValueError("invalid import headers")
        headers = _validate_headers([cell.value for cell in header_cells])
        raw_rows: list[tuple[int, dict]] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"formula cell rejected at row {row_number}")
            values = [cell.value for cell in cells]
            if not any(value is not None for value in values):
                continue
            if len(raw_rows) >= MAX_IMPORT_ROWS:
                raise ValueError(f"import exceeds {MAX_IMPORT_ROWS} row limit")
            raw_rows.append((row_number, dict(zip(headers, values))))
        return _commit_import(conn, table, _prepare_import(conn, table, raw_rows, collision), collision)
    finally:
        book.close()
